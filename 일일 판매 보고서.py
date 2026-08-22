import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.chart import BarChart, Reference

def create_sales_report():
    """
    Simulates a quiz exercise where we create an Excel report,
    format it, and add a chart.
    """
    print("Starting Quiz 54 Solution: Creating Sales Report...")

    # 1. Create a new workbook and select active sheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Data"

    # 2. Add header row
    ws.append(["Date", "Product", "Sales Amount", "Target Achieved"])

    # 3. Add sample data
    data = [
        ["2023-10-01", "Laptop", 1500, "Yes"],
        ["2023-10-02", "Mouse", 25, "No"],
        ["2023-10-03", "Keyboard", 75, "Yes"],
        ["2023-10-04", "Monitor", 300, "Yes"],
        ["2023-10-05", "Mousepad", 10, "No"]
    ]

    for row in data:
        ws.append(row)

    # 4. Format header row
    for cell in ws[1]:
        cell.font = Font(bold=True, color="0000FF") # Blue and Bold
        cell.alignment = Alignment(horizontal="center")

    # 5. Create a Bar Chart for Sales Amount
    chart = BarChart()
    chart.title = "Daily Sales Amount"
    chart.x_axis.title = "Date"
    chart.y_axis.title = "Amount ($)"

    # Define data range for chart (Sales Amount column, excluding header)
    data_ref = Reference(ws, min_col=3, min_row=2, max_col=3, max_row=6)
    
    # Define category range for x-axis (Date column, excluding header)
    cat_ref = Reference(ws, min_col=1, min_row=2, max_col=1, max_row=6)

    chart.add_data(data_ref, titles_from_data=False)
    chart.set_categories(cat_ref)

    # Place chart on the worksheet
    ws.add_chart(chart, "E2")

    # 6. Save the workbook
    filename = "Quiz_54_Report.xlsx"
    wb.save(filename)
    print(f"Report saved successfully as '{filename}'")

if __name__ == "__main__":
    create_sales_report()
